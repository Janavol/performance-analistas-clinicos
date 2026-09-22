"""Geração do relatório em Excel (.xlsx)."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from engine import CAT_LABELS

HEADERS = [
    "Analista", "RC Total", "RC Notificados (n.)", "RC Notificados (%)",
    "RC Notificados no prazo (n.)", "RC Notificados no prazo (%)", "RC Média Min. até notificação",
    "PA Liberados (n.)", "PA Liberado no prazo (n.)", "PA Liberado no prazo (%)",
    "UI Liberados (n.)", "UI Liberado no prazo (n.)", "UI Liberado no prazo (%)",
    "UTI Liberados (n.)", "UTI Liberado no prazo (n.)", "UTI Liberado no prazo (%)",
    "Sepse Liberados (n.)", "Sepse no prazo (n.)", "Sepse no prazo (%)",
    "DT Liberados (n.)", "DT no prazo (n.)", "DT no prazo (%)",
    "AVE Liberados (n.)", "AVE no prazo (n.)", "AVE no prazo (%)",
    "Score Res. Críticos", "Score UCA", "Score Protocolos", "Score PA", "Score UI",
    "Score Final", "Classificação", "Score Final (barra)",
]
PCT_COLS_0IDX = [3, 5, 9, 12, 15, 18, 21, 24]


def _score_bar(v):
    if v is None:
        return ""
    filled = round(max(0, min(100, v)) / 5)  # 20 blocos
    return "█" * filled + "░" * (20 - filled) + f"  {v:.1f}"


def _row(a: dict) -> list:
    return [
        a["analista"], a["rc_total"], a["rc_notif"], a["rc_notif_pct"],
        a["rc_prazo"], a["rc_prazo_pct"], a["rc_media"],
        a["pa"]["n"], a["pa"]["np"], a["pa"]["pct"],
        a["ui"]["n"], a["ui"]["np"], a["ui"]["pct"],
        a["uti"]["n"], a["uti"]["np"], a["uti"]["pct"],
        a["sepse"]["n"], a["sepse"]["np"], a["sepse"]["pct"],
        a["dt"]["n"], a["dt"]["np"], a["dt"]["pct"],
        a["ave"]["n"], a["ave"]["np"], a["ave"]["pct"],
        a["score_criticos"], a["score_uca"], a["score_protocolos"], a["score_pa"], a["score_ui"],
        a["score_final"], a["classe"], _score_bar(a["score_final"]),
    ]


def build_excel(result: dict, config: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Performance"

    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    team_row = dict(result["team"])
    ws.append(_row(team_row))
    for a in result["per_analyst"]:
        ws.append(_row(a))

    for r in range(2, ws.max_row + 1):
        for c in PCT_COLS_0IDX:
            cell = ws.cell(row=r, column=c + 1)
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.0%"

    for i, h in enumerate(HEADERS):
        col_letter = get_column_letter(i + 1)
        ws.column_dimensions[col_letter].width = 30 if i == len(HEADERS) - 1 else max(12, min(28, len(h) + 2))

    ws_cfg = wb.create_sheet("Configuração usada")
    rows = [
        ["Configuração usada nesta análise"], [],
        ["Limites de TAT (minutos)"],
        *[[CAT_LABELS.get(k, k), v] for k, v in config["tat"].items()], [],
        ["Peso dos blocos no score final"],
        ["Resultados críticos", config["block_weights"]["criticos"]],
        ["Protocolos", config["block_weights"]["protocolos"]],
        ["Pronto atendimento", config["block_weights"]["pa"]],
        ["Unidades de Internação", config["block_weights"]["ui"]],
        ["UTI Adulto", config["block_weights"]["uti"]], [],
        ["Composição — resultados críticos"],
        ["Peso taxa de notificação", config["criticos"]["taxa_notificacao"]],
        ["Peso taxa dentro do prazo", config["criticos"]["taxa_prazo"]],
        ["Penalidade por não notificado", config["criticos"]["penalidade"]], [],
        ["Peso interno dos protocolos"],
        ["Sepse", config["protocolos"]["SEPSE"]],
        ["Dor torácica", config["protocolos"]["DOR_TORACICA"]],
        ["AVE", config["protocolos"]["AVE"]], [],
        ["Faixas de classificação (score mínimo)"],
        ["Excelente", config["class_min"]["excelente"]],
        ["Bom", config["class_min"]["bom"]],
        ["Regular", config["class_min"]["regular"]],
    ]
    for row in rows:
        ws_cfg.append(row)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def build_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Dados"
    ws.append([
        "Analista", "Categoria", "Hora início (de acordo com a categoria)",
        "Hora Liberação Clínica", "Hora Notificação (só RES. CRITICOS)",
    ])
    ws.append(["ANALISTA_EXEMPLO", "PA", "2026-08-01 08:00", "2026-08-01 08:45", ""])
    ws.append(["ANALISTA_EXEMPLO", "UI", "2026-08-01 08:30", "2026-08-01 09:20", ""])
    ws.append(["ANALISTA_EXEMPLO", "CRITICOAMB", "2026-08-01 09:10", "2026-08-01 09:55", "2026-08-01 10:20"])
    ws.append(["ANALISTA_EXEMPLO", "SEPSE", "2026-08-01 10:00", "2026-08-01 10:04", ""])

    ws_info = wb.create_sheet("Instruções")
    ws_info.append(["Categorias válidas"])
    for cat in ["CRITICOAMB", "CRITICOHOSP", "PA", "UI", "UTI", "SEPSE", "DOR_TORACICA", "AVE"]:
        ws_info.append([cat])
    ws_info.append([])
    ws_info.append(["UI = Unidades de Internação (limite de 60 min, calculado igual ao Pronto Atendimento)."])
    ws_info.append(["Preencha as datas/horas no formato AAAA-MM-DD HH:MM."])
    ws_info.append(["A coluna de notificação só é usada para CRITICOAMB e CRITICOHOSP — deixe em branco nas demais categorias."])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
